"""Warplan deal model: builds a formula-driven Excel workbook from a deal structure.

Runs inside the workspace's Deal Lab microVM:  python model.py job.json out.xlsx
Every number the buyer might change lives on the Inputs sheet; everything else is a live formula, so the
workbook keeps working when a banker, accountant or seller edits it. Mirrors public/js/deal.js (dealModel).
"""
import json
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

job = json.load(open(sys.argv[1]))
out = sys.argv[2]
d = job["deal"]
cur = {"$": "$", "A$": "A$", "C$": "C$", "£": "£", "€": "€"}.get(d.get("cur", "$"), "$")
MONEY = f'"{cur}"#,##0;[Red]-"{cur}"#,##0'
PCT = "0.0%"
X = '0.00"x"'

ink = Font(name="Calibri", size=11)
bold = Font(name="Calibri", size=11, bold=True)
title = Font(name="Calibri", size=16, bold=True)
muted = Font(name="Calibri", size=9, color="6E6E73")
inp = Font(name="Calibri", size=11, color="0A4BCF")  # blue = an input you can change
head_fill = PatternFill("solid", fgColor="F2F2F5")
input_fill = PatternFill("solid", fgColor="EEF4FF")
line = Border(bottom=Side(style="thin", color="D2D2D7"))

wb = Workbook()

# ------------------------------------------------------------------ Inputs
ws = wb.active
ws.title = "Inputs"
ws["A1"] = f"{job.get('company') or 'Target'}: deal model"
ws["A1"].font = title
ws["A2"] = "Blue cells are inputs. Change them and every sheet updates. Built by Warplan."
ws["A2"].font = muted
ws.column_dimensions["A"].width = 38
ws.column_dimensions["B"].width = 18
ws.column_dimensions["C"].width = 48

names = {}


def put(row, label, value, fmt=None, note="", key=None, is_input=True):
    ws.cell(row=row, column=1, value=label).font = ink
    c = ws.cell(row=row, column=2, value=value)
    c.font = inp if is_input else bold
    if is_input:
        c.fill = input_fill
    if fmt:
        c.number_format = fmt
    if note:
        ws.cell(row=row, column=3, value=note).font = muted
    if key:
        names[key] = f"Inputs!$B${row}"
    return f"Inputs!$B${row}"


def section(row, text):
    c = ws.cell(row=row, column=1, value=text)
    c.font = bold
    for col in (1, 2, 3):
        ws.cell(row=row, column=col).fill = head_fill


r = 4
section(r, "The business"); r += 1
put(r, "Revenue (if known)", job.get("revenue") or None, MONEY, "Optional: for margin only", "revenue"); r += 1
put(r, "Reported EBITDA", d["ebitda"], MONEY, "Before add-backs", "ebitda_reported"); r += 1
addbacks = job.get("addbacks") or []
ab_rows = []
for a in addbacks[:8]:
    ab_rows.append(put(r, f"Add-back: {a.get('label', 'Adjustment')}", a.get("amount", 0), MONEY, a.get("note", ""))); r += 1
if not addbacks:
    ab_rows.append(put(r, "Add-back: owner salary above market", 0, MONEY, "Normalise the owner's pay, one-offs, personal expenses")); r += 1
put(r, "Adjusted EBITDA", f"={names['ebitda_reported']}+{'+'.join(ab_rows)}", MONEY, "What the price is based on", "ebitda", is_input=False); r += 1
put(r, "Free cash flow (% of EBITDA)", d["fcfPct"] / 100, PCT, "After tax, capex and working capital, before debt", "fcf_pct"); r += 1
r += 1
section(r, "Price"); r += 1
put(r, "Multiple of EBITDA", d["multiple"], X, "", "multiple"); r += 1
put(r, "Purchase price", f"={names['ebitda']}*{names['multiple']}", MONEY, "", "price", is_input=False); r += 1
r += 1
section(r, "Capital stack (% of price, must add to 100%)"); r += 1
put(r, "Vendor finance (seller note)", d["vf"]["pct"] / 100, PCT, "", "vf_pct"); r += 1
put(r, "  term (years)", d["vf"]["years"], "0.0", "", "vf_years"); r += 1
put(r, "  interest rate", d["vf"]["rate"] / 100, PCT, "", "vf_rate"); r += 1
put(r, "  payment holiday (months)", d["vf"]["holiday"], "0", "No payments and no interest", "vf_holiday"); r += 1
put(r, "  interest-only (months)", d["vf"]["io"], "0", "After the holiday", "vf_io"); r += 1
put(r, "Commercial debt (bank)", d["bank"]["pct"] / 100, PCT, "", "bank_pct"); r += 1
put(r, "  term (years)", d["bank"]["years"], "0.0", "", "bank_years"); r += 1
put(r, "  interest rate", d["bank"]["rate"] / 100, PCT, "", "bank_rate"); r += 1
put(r, "Seller rollover (equity)", d["roll"]["pct"] / 100, PCT, "Seller keeps this % of the shares", "roll_pct"); r += 1
put(r, "Investor capital", d["inv"]["pct"] / 100, PCT, "", "inv_pct"); r += 1
put(r, "  investors' equity stake", d["inv"]["stake"] / 100, PCT, "", "inv_stake"); r += 1
put(r, "Your own cash", d["own"]["pct"] / 100, PCT, "", "own_pct"); r += 1
put(r, "Total allocated", f"={names['vf_pct']}+{names['bank_pct']}+{names['roll_pct']}+{names['inv_pct']}+{names['own_pct']}", PCT, "Must be 100%", "allocated", is_input=False); r += 1
put(r, "Rollover/investor shares are non-voting", "Yes" if d.get("nonVoting", True) else "No", None, "Yes or No", "nonvoting"); r += 1
put(r, "Minimum DSCR the bank wants", 1.5, X, "Debt service coverage", "dscr_min"); r += 1

# ------------------------------------------------------------------ Debt (monthly)
N = names
horizon = max(12, round(max(d["vf"]["years"], d["bank"]["years"], 1) * 12))
dt = wb.create_sheet("Debt")
heads = ["Month", "Year", "Seller note: opening", "interest", "payment", "closing", "Bank: opening", "interest", "payment", "closing", "Total debt service"]
for i, h in enumerate(heads, 1):
    c = dt.cell(row=1, column=i, value=h)
    c.font = bold
    c.fill = head_fill
    c.alignment = Alignment(wrap_text=True)
    dt.column_dimensions[get_column_letter(i)].width = 14 if i > 2 else 8
dt.freeze_panes = "C2"
vf_amt = f"({N['price']}*{N['vf_pct']})"
bank_amt = f"({N['price']}*{N['bank_pct']})"
vf_r = f"({N['vf_rate']}/12)"
bk_r = f"({N['bank_rate']}/12)"
vf_n = f"ROUND({N['vf_years']}*12,0)"
bk_n = f"ROUND({N['bank_years']}*12,0)"
vf_am = f"MAX(1,{vf_n}-{N['vf_holiday']}-{N['vf_io']})"
# Level payment once amortisation starts (same as the app's loanSchedule).
vf_pay = f"IF({vf_r}=0,{vf_amt}/{vf_am},{vf_amt}*{vf_r}/(1-(1+{vf_r})^(-{vf_am})))"
bk_pay = f"IF({bk_r}=0,{bank_amt}/{bk_n},{bank_amt}*{bk_r}/(1-(1+{bk_r})^(-{bk_n})))"
for m in range(1, horizon + 1):
    row = m + 1
    dt.cell(row=row, column=1, value=m)
    dt.cell(row=row, column=2, value=f"=ROUNDUP(A{row}/12,0)")
    dt.cell(row=row, column=3, value=f"={vf_amt}" if m == 1 else f"=F{row - 1}")
    dt.cell(row=row, column=4, value=f"=IF(OR(C{row}<=0.5,A{row}>{vf_n},A{row}<={N['vf_holiday']}),0,C{row}*{vf_r})")
    dt.cell(row=row, column=5, value=(f"=IF(OR(C{row}<=0.5,A{row}>{vf_n}),0,IF(A{row}={vf_n},C{row}+D{row},"
                                      f"IF(A{row}<={N['vf_holiday']},0,IF(A{row}<={N['vf_holiday']}+{N['vf_io']},D{row},MIN({vf_pay},C{row}+D{row})))))"))
    dt.cell(row=row, column=6, value=f"=MAX(0,C{row}+D{row}-E{row})")
    dt.cell(row=row, column=7, value=f"={bank_amt}" if m == 1 else f"=J{row - 1}")
    dt.cell(row=row, column=8, value=f"=IF(OR(G{row}<=0.5,A{row}>{bk_n}),0,G{row}*{bk_r})")
    dt.cell(row=row, column=9, value=f"=IF(OR(G{row}<=0.5,A{row}>{bk_n}),0,IF(A{row}={bk_n},G{row}+H{row},MIN({bk_pay},G{row}+H{row})))")
    dt.cell(row=row, column=10, value=f"=MAX(0,G{row}+H{row}-I{row})")
    dt.cell(row=row, column=11, value=f"=E{row}+I{row}")
    for col in range(3, 12):
        dt.cell(row=row, column=col).number_format = MONEY
last = horizon + 1

# ------------------------------------------------------------------ Annual + DSCR
an = wb.create_sheet("Annual", 1)
an["A1"] = "Year by year"
an["A1"].font = title
heads = ["Year", "Free cash flow", "Seller note service", "Bank service", "Total debt service", "DSCR", "Debt left at year end", "Cash to you after debt"]
for i, h in enumerate(heads, 1):
    c = an.cell(row=3, column=i, value=h)
    c.font = bold
    c.fill = head_fill
    c.alignment = Alignment(wrap_text=True)
    an.column_dimensions[get_column_letter(i)].width = 16
years = horizon // 12
for y in range(1, years + 1):
    row = y + 3
    an.cell(row=row, column=1, value=y)
    an.cell(row=row, column=2, value=f"={N['ebitda']}*{N['fcf_pct']}").number_format = MONEY
    an.cell(row=row, column=3, value=f"=SUMIFS(Debt!$E$2:$E${last},Debt!$B$2:$B${last},A{row})").number_format = MONEY
    an.cell(row=row, column=4, value=f"=SUMIFS(Debt!$I$2:$I${last},Debt!$B$2:$B${last},A{row})").number_format = MONEY
    an.cell(row=row, column=5, value=f"=C{row}+D{row}").number_format = MONEY
    an.cell(row=row, column=6, value=f'=IF(E{row}>0,B{row}/E{row},"")').number_format = X
    an.cell(row=row, column=7, value=f"=INDEX(Debt!$F$2:$F${last},A{row}*12)+INDEX(Debt!$J$2:$J${last},A{row}*12)").number_format = MONEY
    an.cell(row=row, column=8, value=f"=B{row}-E{row}").number_format = MONEY
first, end = 4, years + 3
an.conditional_formatting.add(f"F{first}:F{end}", CellIsRule(operator="lessThan", formula=[N["dscr_min"]], font=Font(color="C0392B", bold=True)))

# ------------------------------------------------------------------ Summary (first sheet people see)
sm = wb.create_sheet("Summary", 0)
sm.column_dimensions["A"].width = 36
sm.column_dimensions["B"].width = 20
sm.column_dimensions["C"].width = 54
sm["A1"] = f"{job.get('company') or 'Target'}"
sm["A1"].font = title
sm["A2"] = job.get("subtitle") or "Deal model"
sm["A2"].font = muted
rows = [
    ("Adjusted EBITDA", f"={N['ebitda']}", MONEY, ""),
    ("Multiple", f"={N['multiple']}", X, ""),
    ("Purchase price", f"={N['price']}", MONEY, ""),
    ("", None, None, ""),
    ("Sources", None, None, ""),
    ("  Seller note", f"={N['price']}*{N['vf_pct']}", MONEY, ""),
    ("  Bank", f"={N['price']}*{N['bank_pct']}", MONEY, ""),
    ("  Seller rollover", f"={N['price']}*{N['roll_pct']}", MONEY, ""),
    ("  Investors", f"={N['price']}*{N['inv_pct']}", MONEY, ""),
    ("  Your cash", f"={N['price']}*{N['own_pct']}", MONEY, ""),
    ("Cash needed at closing", f"={N['price']}*({N['bank_pct']}+{N['inv_pct']}+{N['own_pct']})", MONEY, "Bank + investors + you: what lands with the seller on day one"),
    ("", None, None, ""),
    ("Lowest DSCR in any year", f'=IF(COUNT(Annual!F{first}:F{end})=0,"",MIN(Annual!F{first}:F{end}))', X, "Banks and sellers want 1.5x or more"),
    ("Your economic ownership", f"=MAX(0,1-{N['roll_pct']}-{N['inv_stake']})", PCT, ""),
    ("Your votes", f'=IF({N["nonvoting"]}="Yes",1,MAX(0,1-{N["roll_pct"]}-{N["inv_stake"]}))', PCT, "Over 50% = you control the company"),
    ("Stack adds to 100%", f'=IF(ABS({N["allocated"]}-1)<0.0001,"Yes","No: fix Inputs")', None, ""),
    ("Does the deal work?", f'=IF(AND(B19="Yes",B18>0.5,OR(B16="",B16>={N["dscr_min"]})),"Yes","Not yet")', None, "Complete stack, you control it, DSCR ≥ the bank's minimum every year"),
]
for i, (label, value, fmt, note) in enumerate(rows, start=4):
    sm.cell(row=i, column=1, value=label).font = bold if label and not label.startswith("  ") else ink
    if value is not None:
        c = sm.cell(row=i, column=2, value=value)
        c.font = bold
        if fmt:
            c.number_format = fmt
    if note:
        sm.cell(row=i, column=3, value=note).font = muted
    sm.cell(row=i, column=1).border = line
    sm.cell(row=i, column=2).border = line

# Sensitivity: price at different multiples and EBITDA levels (plain formulas, no data tables needed).
sm["A24"] = "Price sensitivity"
sm["A24"].font = bold
mults = [max(0.5, round(d["multiple"] + k * 0.5, 2)) for k in (-2, -1, 0, 1, 2)]
swings = [-0.2, -0.1, 0, 0.1, 0.2]
sm.cell(row=25, column=1, value="EBITDA change ↓ / multiple →").font = muted
for j, mu in enumerate(mults):
    c = sm.cell(row=25, column=2 + j, value=mu)
    c.number_format = X
    c.font = bold
    sm.column_dimensions[get_column_letter(2 + j)].width = max(sm.column_dimensions[get_column_letter(2 + j)].width or 0, 14)
for i, sw in enumerate(swings):
    c = sm.cell(row=26 + i, column=1, value=sw)
    c.number_format = '+0%;-0%;0%'
    for j in range(len(mults)):
        col = get_column_letter(2 + j)
        sm.cell(row=26 + i, column=2 + j, value=f"={N['ebitda']}*(1+$A{26 + i})*{col}$25").number_format = MONEY

if job.get("notes"):
    sm["A33"] = "Notes"
    sm["A33"].font = bold
    sm["A34"] = job["notes"]
    sm["A34"].alignment = Alignment(wrap_text=True, vertical="top")
    sm.merge_cells("A34:F40")

for sheet in wb.worksheets:
    sheet.sheet_view.showGridLines = False
wb.active = 0
wb.save(out)
print(json.dumps({"ok": True, "months": horizon, "years": years}))
